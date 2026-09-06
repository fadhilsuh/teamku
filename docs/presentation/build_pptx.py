# /// script
# requires-python = ">=3.10"
# dependencies = ["python-pptx>=1.0.0"]
# ///
"""Generate the Teamku benefit-led sales deck."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt


EMU = 914400


def inch(value: float) -> Emu:
    return Emu(int(value * EMU))


# PowerPoint-safe solid colors. No custom OOXML gradients.
NAVY = RGBColor(10, 16, 38)
NAVY_2 = RGBColor(18, 27, 58)
INK = RGBColor(243, 246, 255)
MUTED = RGBColor(166, 177, 205)
PURPLE = RGBColor(124, 92, 252)
CYAN = RGBColor(0, 212, 199)
CORAL = RGBColor(255, 107, 107)
AMBER = RGBColor(255, 184, 77)
LIME = RGBColor(163, 230, 53)
CREAM = RGBColor(247, 245, 239)
DARK_TEXT = RGBColor(24, 31, 55)
LIGHT_MUTED = RGBColor(91, 101, 126)
WHITE = RGBColor(255, 255, 255)
FONT = "Aptos"


prs = Presentation()
prs.slide_width = inch(13.333)
prs.slide_height = inch(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]
MX = inch(0.72)
CW = SW - 2 * MX


def shape(slide, x, y, w, h, color, radius=True, line=None):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        x, y, w, h,
    )
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    if line:
        shp.line.color.rgb = line
        shp.line.width = Pt(1)
    else:
        shp.line.fill.background()
    if radius:
        try:
            shp.adjustments[0] = 0.08
        except Exception:
            pass
    return shp


def add_text(
    slide, x, y, w, h, paragraphs, *,
    align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
    margins=0.03, line_spacing=1.0,
):
    """paragraphs: list[list[(text, size, color, bold)]]."""
    box = slide.shapes.add_textbox(x, y, w, h)
    frame = box.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = anchor
    frame.margin_left = frame.margin_right = inch(margins)
    frame.margin_top = frame.margin_bottom = inch(margins)
    for idx, runs in enumerate(paragraphs):
        p = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(3)
        p.line_spacing = line_spacing
        for content, size, color, bold in runs:
            r = p.add_run()
            r.text = content
            r.font.name = FONT
            r.font.size = Pt(size)
            r.font.color.rgb = color
            r.font.bold = bold
    return box


def dark_slide(accent=PURPLE):
    slide = prs.slides.add_slide(BLANK)
    shape(slide, 0, 0, SW, SH, NAVY, radius=False)
    # Decorative shapes create depth while staying PowerPoint-safe.
    shape(slide, SW - inch(2.4), -inch(1.0), inch(3.2), inch(3.2), accent)
    shape(slide, SW - inch(1.7), -inch(0.35), inch(2.0), inch(2.0), NAVY_2)
    shape(slide, -inch(0.9), SH - inch(1.0), inch(1.7), inch(1.7), NAVY_2)
    return slide


def light_slide(accent=PURPLE):
    slide = prs.slides.add_slide(BLANK)
    shape(slide, 0, 0, SW, SH, CREAM, radius=False)
    shape(slide, 0, 0, inch(0.18), SH, accent, radius=False)
    shape(slide, SW - inch(1.4), -inch(0.65), inch(2.0), inch(2.0), accent)
    return slide


def kicker(slide, label, color=CYAN, dark=True):
    add_text(
        slide, MX, inch(0.48), CW, inch(0.35),
        [[(label.upper(), 11, color, True)]],
    )


def heading(slide, title, y=0.92, dark=True, size=34, w=None):
    add_text(
        slide, MX, inch(y), w or CW, inch(1.05),
        [[(title, size, INK if dark else DARK_TEXT, True)]],
        line_spacing=0.95,
    )


def body(slide, content, x=MX, y=1.85, w=None, h=0.8, dark=True, size=14):
    add_text(
        slide, x, inch(y), w or CW, inch(h),
        [[(content, size, MUTED if dark else LIGHT_MUTED, False)]],
        line_spacing=1.15,
    )


def icon_badge(slide, x, y, label, color, size=0.42, dark_text=False):
    badge = shape(slide, x, y, inch(size), inch(size), color)
    badge.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    badge.text_frame.margin_left = badge.text_frame.margin_right = 0
    badge.text_frame.margin_top = badge.text_frame.margin_bottom = 0
    p = badge.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.name = FONT
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = DARK_TEXT if dark_text else WHITE
    return badge


def card(slide, x, y, w, h, code, title, description, accent=CYAN, light=False):
    bg = WHITE if light else NAVY_2
    title_color = DARK_TEXT if light else INK
    desc_color = LIGHT_MUTED if light else MUTED
    shape(slide, x, y, w, h, bg, line=RGBColor(225, 226, 231) if light else RGBColor(38, 49, 85))
    icon_badge(slide, x + inch(0.2), y + inch(0.18), code, accent, dark_text=accent in (CYAN, AMBER, LIME))
    add_text(
        slide, x + inch(0.75), y + inch(0.18), w - inch(0.95), inch(0.4),
        [[(title, 13, title_color, True)]],
    )
    add_text(
        slide, x + inch(0.2), y + inch(0.78), w - inch(0.4), h - inch(0.92),
        [[(description, 10.5, desc_color, False)]],
        line_spacing=1.12,
    )


def metric(slide, x, y, w, value, label, accent):
    shape(slide, x, y, w, inch(1.35), NAVY_2, line=RGBColor(38, 49, 85))
    shape(slide, x, y, inch(0.08), inch(1.35), accent, radius=False)
    add_text(slide, x + inch(0.22), y + inch(0.16), w - inch(0.35), inch(0.55),
             [[(value, 27, accent, True)]])
    add_text(slide, x + inch(0.22), y + inch(0.78), w - inch(0.35), inch(0.4),
             [[(label, 10.5, MUTED, False)]])


def footer(slide, number, dark=True):
    color = MUTED if dark else LIGHT_MUTED
    add_text(slide, MX, SH - inch(0.45), inch(5), inch(0.25),
             [[("TEAMKU  ·  MOVON DIGITAL HOUSE", 8, color, True)]])
    add_text(slide, SW - MX - inch(1.2), SH - inch(0.45), inch(1.2), inch(0.25),
             [[(f"{number:02d} / 10", 8, color, True)]], align=PP_ALIGN.RIGHT)


def pill(slide, x, y, label, color):
    width = inch(0.34 + len(label) * 0.085)
    p = shape(slide, x, y, width, inch(0.42), color)
    p.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    p.text_frame.margin_left = p.text_frame.margin_right = 0
    p.text_frame.margin_top = p.text_frame.margin_bottom = 0
    para = p.text_frame.paragraphs[0]
    para.alignment = PP_ALIGN.CENTER
    run = para.add_run()
    run.text = label
    run.font.name = FONT
    run.font.size = Pt(10)
    run.font.bold = True
    run.font.color.rgb = DARK_TEXT if color in (CYAN, AMBER, LIME) else WHITE
    return width


def feature_row(slide, x, y, w, code, title, description, accent):
    icon_badge(slide, x, y + inch(0.04), code, accent, size=0.48, dark_text=accent in (CYAN, AMBER, LIME))
    add_text(slide, x + inch(0.68), y, w - inch(0.68), inch(0.35),
             [[(title, 13, INK, True)]])
    add_text(slide, x + inch(0.68), y + inch(0.37), w - inch(0.68), inch(0.58),
             [[(description, 10.5, MUTED, False)]], line_spacing=1.08)


# 1 — Cover
s = dark_slide(PURPLE)
shape(s, inch(8.55), inch(1.1), inch(3.7), inch(5.2), PURPLE)
shape(s, inch(8.85), inch(1.4), inch(3.1), inch(4.6), NAVY_2)
icon_badge(s, inch(9.3), inch(1.85), "01", CYAN, size=0.62, dark_text=True)
add_text(s, inch(9.3), inch(2.65), inch(2.45), inch(2.6), [
    [("Satu platform.", 24, INK, True)],
    [("Semua urusan HR.", 24, INK, True)],
    [("Lebih sedikit administrasi.", 24, CYAN, True)],
])
add_text(s, inch(9.3), inch(5.2), inch(2.2), inch(0.45),
         [[("PEOPLE OS", 10, MUTED, True)]])

kicker(s, "Platform HR Indonesia")
add_text(s, MX, inch(1.45), inch(7.1), inch(1.2), [[("Teamku", 58, INK, True)]])
add_text(s, MX, inch(2.75), inch(7.2), inch(1.75), [
    [("Kerja HR lebih ringan.", 31, INK, True)],
    [("Tim bergerak lebih cepat.", 31, CYAN, True)],
], line_spacing=0.95)
body(
    s,
    "Kehadiran, cuti, persetujuan, dan payroll bekerja dalam satu alur—"
    "agar HR bisa fokus membangun tim, bukan mengejar administrasi.",
    y=4.75, w=inch(6.9), h=0.9, size=15,
)
x = MX
for label, color in [("Hemat waktu", PURPLE), ("Data tepercaya", CYAN), ("Tim lebih mandiri", CORAL)]:
    x += pill(s, x, inch(5.95), label, color) + inch(0.14)
footer(s, 1)

# 2 — Before / after
s = light_slide(CORAL)
kicker(s, "Perubahan yang terasa", CORAL, dark=False)
heading(s, "Dari administrasi yang tercecer menjadi satu alur yang terkendali", dark=False, size=31)
body(s, "Teamku merapikan perjalanan kerja HR dari awal sampai akhir—tanpa berpindah-pindah chat, spreadsheet, dan dokumen.", y=1.83, h=0.65, dark=False)

gap = inch(0.35)
col = (CW - gap) / 2
left_x = MX
right_x = MX + col + gap
shape(s, left_x, inch(2.65), col, inch(3.75), WHITE, line=RGBColor(225, 226, 231))
shape(s, right_x, inch(2.65), col, inch(3.75), NAVY)
add_text(s, left_x + inch(0.28), inch(2.9), col - inch(0.56), inch(0.4),
         [[("SEBELUM", 11, CORAL, True)]])
add_text(s, right_x + inch(0.28), inch(2.9), col - inch(0.56), inch(0.4),
         [[("BERSAMA TEAMKU", 11, CYAN, True)]])
before = [
    ("01", "Absensi sulit diverifikasi"),
    ("02", "Cuti terselip di percakapan"),
    ("03", "Payroll penuh pengecekan ulang"),
    ("04", "Data tersebar, keputusan terlambat"),
]
after = [
    ("01", "Kehadiran dilengkapi bukti lokasi & selfie"),
    ("02", "Persetujuan masuk ke satu inbox"),
    ("03", "Payroll mengalir dari susun hingga publikasi"),
    ("04", "Dashboard memberi visibilitas real-time"),
]
for idx, (code, label) in enumerate(before):
    y = inch(3.48 + idx * 0.67)
    icon_badge(s, left_x + inch(0.28), y, code, RGBColor(239, 231, 229), dark_text=True)
    add_text(s, left_x + inch(0.82), y + inch(0.03), col - inch(1.1), inch(0.38),
             [[(label, 11.5, DARK_TEXT, True)]])
for idx, (code, label) in enumerate(after):
    y = inch(3.48 + idx * 0.67)
    icon_badge(s, right_x + inch(0.28), y, code, CYAN, dark_text=True)
    add_text(s, right_x + inch(0.82), y + inch(0.03), col - inch(1.1), inch(0.38),
             [[(label, 11.5, INK, True)]])
footer(s, 2, dark=False)

# 3 — Benefits
s = dark_slide(CYAN)
kicker(s, "Dampak untuk bisnis")
heading(s, "Bukan sekadar digital. Cara kerja tim Anda ikut berubah.")
body(s, "Empat manfaat yang langsung terasa ketika seluruh proses HR bergerak dalam satu sistem.", y=1.82, h=0.55)
benefits = [
    ("01", "Waktu kembali ke tim", "Kurangi pekerjaan berulang agar HR punya ruang untuk mendampingi karyawan.", PURPLE),
    ("02", "Data lebih tepercaya", "Kehadiran dan aktivitas penting dilengkapi jejak yang mudah ditelusuri.", CYAN),
    ("03", "Keputusan lebih cepat", "Manager melihat yang perlu diputuskan tanpa menunggu rekap manual.", CORAL),
    ("04", "Pengalaman lebih rapi", "Karyawan mengurus kebutuhan HR melalui alur mandiri yang konsisten.", AMBER),
]
card_w = (CW - inch(0.54)) / 2
for idx, (code, name, desc, accent) in enumerate(benefits):
    row, col_idx = divmod(idx, 2)
    x = MX + col_idx * (card_w + inch(0.54))
    y = inch(2.65 + row * 1.75)
    card(s, x, y, card_w, inch(1.48), code, name, desc, accent)
footer(s, 3)

# 4 — Attendance
s = dark_slide(CYAN)
kicker(s, "Kehadiran")
add_text(s, MX, inch(1.18), inch(6.4), inch(1.85), [
    [("Datang.", 40, INK, True)],
    [("Buktikan.", 40, CYAN, True)],
    [("Mulai bekerja.", 40, INK, True)],
], line_spacing=0.88)
body(
    s,
    "Presensi menjadi bagian dari ritme kerja—bukan sekadar tombol check-in. "
    "Setiap kehadiran disertai konteks lokasi, bukti visual, dan agenda hari itu.",
    y=3.45, w=inch(5.9), h=1.25, size=15,
)
shape(s, MX, inch(5.15), inch(5.65), inch(0.08), CYAN, radius=False)
body(s, "Benefit: mengurangi risiko titip absen dan memberi visibilitas awal atas fokus tim.", y=5.43, w=inch(5.9), h=0.75, size=12)

panel_x = inch(7.35)
shape(s, panel_x, inch(1.15), inch(4.75), inch(5.25), NAVY_2, line=RGBColor(38, 49, 85))
add_text(s, panel_x + inch(0.35), inch(1.5), inch(4.0), inch(0.4),
         [[("ALUR KEHADIRAN", 10, CYAN, True)]])
steps = [
    ("01", "Ambil selfie langsung", "Bukti visual diambil saat check-in.", PURPLE),
    ("02", "Verifikasi geofence", "Lokasi diperiksa terhadap radius kantor.", CYAN),
    ("03", "Tetapkan agenda", "Karyawan menuliskan 1–5 prioritas.", CORAL),
    ("04", "Tutup dengan ringkasan", "Check-out mencatat hasil kerja hari itu.", AMBER),
]
for idx, (code, name, desc, accent) in enumerate(steps):
    y = inch(2.05 + idx * 1.0)
    feature_row(s, panel_x + inch(0.35), y, inch(4.0), code, name, desc, accent)
footer(s, 4)

# 5 — Leave
s = light_slide(PURPLE)
kicker(s, "Cuti & Persetujuan", PURPLE, dark=False)
heading(s, "Cuti tanpa mengejar chat. Keputusan tanpa menunggu rekap.", dark=False)
body(s, "Satu alur yang jelas bagi karyawan, manager, dan HR—dari pengajuan sampai notifikasi keputusan.", y=1.82, h=0.6, dark=False)

timeline_y = inch(3.15)
line_y = timeline_y + inch(0.34)
shape(s, MX + inch(0.25), line_y, CW - inch(0.5), inch(0.07), RGBColor(218, 216, 229), radius=False)
steps = [
    ("1", "Ajukan", "Pilih tanggal & alasan", PURPLE),
    ("2", "Hitung", "Hari kerja & saldo otomatis", CYAN),
    ("3", "Putuskan", "Approve/reject + catatan", CORAL),
    ("4", "Ketahui", "Notifikasi diterima seketika", AMBER),
]
step_w = CW / 4
for idx, (code, name, desc, accent) in enumerate(steps):
    x = MX + idx * step_w
    icon_badge(s, x + step_w / 2 - inch(0.34), timeline_y, code, accent, size=0.68, dark_text=accent in (CYAN, AMBER))
    add_text(s, x, inch(4.02), step_w, inch(0.35), [[(name, 14, DARK_TEXT, True)]], align=PP_ALIGN.CENTER)
    add_text(s, x + inch(0.12), inch(4.48), step_w - inch(0.24), inch(0.6),
             [[(desc, 10.5, LIGHT_MUTED, False)]], align=PP_ALIGN.CENTER, line_spacing=1.08)
shape(s, MX, inch(5.45), CW, inch(0.9), NAVY)
add_text(s, MX + inch(0.3), inch(5.66), inch(8.5), inch(0.4),
         [[("Hasilnya:", 12, CYAN, True), ("  status transparan, saldo selalu jelas, dan approval tidak lagi tercecer.", 12, INK, True)]])
footer(s, 5, dark=False)

# 6 — Payroll
s = dark_slide(PURPLE)
kicker(s, "Payroll")
heading(s, "Payroll yang terasa tenang—bahkan di akhir bulan")
body(s, "Angka tersusun jelas, perubahan terkendali, dan setiap karyawan menerima slip gajinya secara privat.", y=1.82, h=0.65)

labels = [
    ("01", "SUSUN", "Siapkan payroll bulanan dan tinjau komponen per karyawan.", PURPLE),
    ("02", "KUNCI", "Finalisasi angka agar tidak berubah tanpa sengaja.", CORAL),
    ("03", "PUBLIKASIKAN", "Terbitkan slip dan beri notifikasi kepada seluruh tim.", CYAN),
]
box_w = (CW - inch(0.5)) / 3
for idx, (code, name, desc, accent) in enumerate(labels):
    x = MX + idx * (box_w + inch(0.25))
    shape(s, x, inch(2.75), box_w, inch(2.2), NAVY_2, line=accent)
    icon_badge(s, x + inch(0.25), inch(3.0), code, accent, size=0.55, dark_text=accent in (CYAN, AMBER))
    add_text(s, x + inch(0.25), inch(3.72), box_w - inch(0.5), inch(0.4),
             [[(name, 15, INK, True)]])
    add_text(s, x + inch(0.25), inch(4.18), box_w - inch(0.5), inch(0.58),
             [[(desc, 10.5, MUTED, False)]], line_spacing=1.08)
metric(s, MX, inch(5.35), inch(3.5), "IDR", "Kalkulasi presisi tanpa salah pembulatan", AMBER)
metric(s, MX + inch(3.75), inch(5.35), inch(3.5), "PRIVAT", "Setiap karyawan hanya melihat slipnya", CYAN)
metric(s, MX + inch(7.5), inch(5.35), inch(3.5), "TERCATAT", "Setiap tahap memiliki jejak audit", PURPLE)
footer(s, 6)

# 7 — Employee experience
s = light_slide(CYAN)
kicker(s, "Pengalaman Karyawan", CYAN, dark=False)
heading(s, "HR tanpa antrean: semua yang dibutuhkan ada di tangan karyawan", dark=False, size=32)
body(s, "Pengalaman mandiri yang sederhana mengurangi pertanyaan berulang—dan membuat layanan HR terasa cepat.", y=1.82, h=0.6, dark=False)

phone_x = inch(1.1)
shape(s, phone_x, inch(2.6), inch(3.0), inch(3.75), NAVY)
shape(s, phone_x + inch(0.18), inch(2.82), inch(2.64), inch(3.28), NAVY_2)
add_text(s, phone_x + inch(0.42), inch(3.1), inch(2.15), inch(0.4),
         [[("Halo, Tim 👋", 17, INK, True)]])
for idx, (label, color) in enumerate([("Check-in hari ini", CYAN), ("Ajukan cuti", PURPLE), ("Lihat slip gaji", CORAL)]):
    y = inch(3.78 + idx * 0.65)
    shape(s, phone_x + inch(0.42), y, inch(2.15), inch(0.48), color)
    add_text(s, phone_x + inch(0.52), y + inch(0.09), inch(1.95), inch(0.25),
             [[(label, 10, DARK_TEXT if color == CYAN else WHITE, True)]])

rx = inch(5.05)
items = [
    ("01", "Kehadiran hari ini", "Check-in, agenda, dan status check-out terlihat jelas.", CYAN),
    ("02", "Saldo & status cuti", "Tidak perlu bertanya ke HR untuk mengetahui sisa hak cuti.", PURPLE),
    ("03", "Slip gaji pribadi", "Dokumen tersedia aman begitu payroll dipublikasikan.", CORAL),
    ("04", "Notifikasi penting", "Keputusan dan pembaruan masuk ke satu inbox.", AMBER),
]
for idx, item in enumerate(items):
    code, name, desc, accent = item
    y = inch(2.55 + idx * 0.95)
    icon_badge(s, rx, y, code, accent, size=0.48, dark_text=accent in (CYAN, AMBER))
    add_text(s, rx + inch(0.68), y, inch(6.2), inch(0.35), [[(name, 13, DARK_TEXT, True)]])
    add_text(s, rx + inch(0.68), y + inch(0.38), inch(6.2), inch(0.42),
             [[(desc, 10.5, LIGHT_MUTED, False)]])
footer(s, 7, dark=False)

# 8 — Manager dashboard
s = dark_slide(AMBER)
kicker(s, "Visibilitas untuk Manager", AMBER)
heading(s, "Dari data harian menjadi keputusan yang lebih cepat")
body(s, "Satu layar untuk melihat apa yang terjadi hari ini—tanpa menunggu spreadsheet selesai direkap.", y=1.82, h=0.55)

metric(s, MX, inch(2.65), inch(2.75), "18", "Total anggota tim", CYAN)
metric(s, MX + inch(2.98), inch(2.65), inch(2.75), "15", "Hadir hari ini", LIME)
metric(s, MX + inch(5.96), inch(2.65), inch(2.75), "3", "Menunggu persetujuan", CORAL)
metric(s, MX + inch(8.94), inch(2.65), inch(2.75), "83%", "Agenda telah disusun", AMBER)

shape(s, MX, inch(4.35), CW, inch(1.65), NAVY_2, line=RGBColor(38, 49, 85))
add_text(s, MX + inch(0.28), inch(4.62), inch(2.2), inch(0.35),
         [[("YANG LANGSUNG TERLIHAT", 10, AMBER, True)]])
insights = [
    ("Hadir & terlambat", CYAN),
    ("Approval tertunda", CORAL),
    ("Progres agenda", LIME),
    ("Cuti aktif", PURPLE),
]
x = MX + inch(0.3)
for label, color in insights:
    x += pill(s, x, inch(5.18), label, color) + inch(0.17)
footer(s, 8)

# 9 — Roles and modules
s = light_slide(PURPLE)
kicker(s, "Satu platform, tiga pengalaman", PURPLE, dark=False)
heading(s, "Setiap peran melihat yang dibutuhkan—tidak lebih, tidak kurang", dark=False, size=31)
body(s, "Akses yang tepat menjaga privasi sekaligus membuat setiap orang tetap produktif.", y=1.82, h=0.55, dark=False)

role_w = (CW - inch(0.5)) / 3
roles = [
    ("E", "KARYAWAN", "Presensi · agenda · cuti · slip gaji · notifikasi", CYAN),
    ("M", "MANAGER", "Dashboard tim · direktori · persetujuan departemen", PURPLE),
    ("HR", "HR ADMIN", "Data karyawan · payroll · pengaturan · audit log", CORAL),
]
for idx, (code, name, desc, accent) in enumerate(roles):
    x = MX + idx * (role_w + inch(0.25))
    shape(s, x, inch(2.62), role_w, inch(2.3), WHITE, line=RGBColor(225, 226, 231))
    icon_badge(s, x + inch(0.25), inch(2.9), code, accent, size=0.58, dark_text=accent == CYAN)
    add_text(s, x + inch(0.25), inch(3.67), role_w - inch(0.5), inch(0.35),
             [[(name, 13, DARK_TEXT, True)]])
    add_text(s, x + inch(0.25), inch(4.12), role_w - inch(0.5), inch(0.58),
             [[(desc, 10.5, LIGHT_MUTED, False)]], line_spacing=1.08)

add_text(s, MX, inch(5.37), CW, inch(0.35), [[("9 MODUL YANG BEKERJA BERSAMA", 10, PURPLE, True)]])
x = MX
module_colors = [PURPLE, CYAN, CORAL, AMBER, PURPLE, CYAN, CORAL, AMBER, PURPLE]
for label, color in zip(
    ["Beranda", "Karyawan", "Presensi", "Cuti", "Approval", "Payroll", "Slip", "Notifikasi", "Pengaturan"],
    module_colors,
):
    width = pill(s, x, inch(5.9), label, color)
    x += width + inch(0.1)
footer(s, 9, dark=False)

# 10 — CTA
s = dark_slide(CORAL)
shape(s, inch(8.55), inch(0.85), inch(3.6), inch(5.85), CORAL)
shape(s, inch(8.9), inch(1.2), inch(2.9), inch(5.15), NAVY_2)
add_text(s, inch(9.25), inch(1.75), inch(2.25), inch(0.4),
         [[("TEAMKU", 11, CORAL, True)]])
add_text(s, inch(9.25), inch(2.35), inch(2.25), inch(2.4), [
    [("Lebih rapi.", 24, INK, True)],
    [("Lebih cepat.", 24, INK, True)],
    [("Lebih manusiawi.", 24, CYAN, True)],
], line_spacing=1.05)
add_text(s, inch(9.25), inch(5.28), inch(2.2), inch(0.65),
         [[("PEOPLE OS", 10, MUTED, True)], [("MOVON DIGITAL HOUSE", 9, MUTED, False)]])

kicker(s, "Saatnya bergerak")
add_text(s, MX, inch(1.55), inch(7.2), inch(2.1), [
    [("Buat HR bekerja", 40, INK, True)],
    [("secerdas tim Anda.", 40, CYAN, True)],
], line_spacing=0.93)
body(
    s,
    "Satukan kehadiran, cuti, persetujuan, payroll, dan data karyawan "
    "dalam satu pengalaman yang mudah digunakan setiap hari.",
    y=3.95, w=inch(6.8), h=0.95, size=15,
)
x = MX
for label, color in [("Kurangi administrasi", PURPLE), ("Percepat keputusan", CYAN), ("Bangun kepercayaan", CORAL)]:
    x += pill(s, x, inch(5.3), label, color) + inch(0.14)
add_text(s, MX, inch(6.18), inch(5.0), inch(0.45),
         [[("TEAMKU  ·  PEOPLE OS", 13, INK, True)]])
footer(s, 10)


output = Path(__file__).parent / "teamku-deck.pptx"
prs.save(output)
print(f"saved: {output} ({len(prs.slides)} slides)")
